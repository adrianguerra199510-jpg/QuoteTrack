import pytest
from docx import Document
from PySide6.QtWidgets import QFileDialog, QMessageBox

from app import db, documento, ui_main, word
from app.calculo import Cargo, Linea


def abrir(ruta):
    return Document(str(ruta))


def textos(doc):
    """Todo el texto del documento: párrafos y celdas de tablas."""
    partes = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for fila in t.rows:
            partes += [c.text for c in fila.cells]
    return "\n".join(partes)


def tabla_items(doc):
    return next(t for t in doc.tables if t.rows[0].cells[0].text == "Ítem")


@pytest.fixture
def obra(conn):
    db.guardar_cotizacion(1, [
        Linea("grupo", 0, "Albañilería"),
        Linea("item", 1, "Muro de bloque de 6", 10, 20.0, True, "m²"),
        Linea("item", 1, "Muro de bloque de 4", 5, 10.0, False, "m²"),
        Linea("item", 0, "Limpieza", 1, 100.0, True, "gl"),
    ], [Cargo("Administración", "porcentaje", 10)], conn)
    db.guardar_notas(1, "ALCANCE\n• Obra civil según planos.\n- Segunda viñeta\n\nFORMA DE PAGO\n• 50% al inicio.\nTexto suelto", conn)
    return conn


def test_cotizacion_contenido_basico(conn, tmp_path):
    ruta = tmp_path / "c.docx"
    word.exportar_cotizacion(1, str(ruta))
    t = textos(abrir(ruta))
    for esperado in ("INTEGRO SA", "COTIZACIÓN", "COT-2026-0001", "ACME", "8-123-456", "Servicio",
                     "B/. 200.00", "ITBMS (7%):", "B/. 14.00", "B/. 214.00"):
        assert esperado in t, esperado


def test_grupos_codigos_exentos_y_costos(obra, tmp_path):
    ruta = tmp_path / "c.docx"
    word.exportar_cotizacion(1, str(ruta))
    doc = abrir(ruta)
    filas = [[c.text for c in f.cells] for f in tabla_items(doc).rows]
    assert filas[0] == ["Ítem", "Descripción", "Ud.", "Cantidad", "Precio unitario", "Importe"]
    assert [f[0] for f in filas[1:]] == ["01", "01.1", "01.2", "02"]
    assert filas[1][5] == "B/. 250.00" and filas[2][2] == "m²" and filas[2][3] == "10"
    assert "(exento de ITBMS)" in filas[3][1] and "(exento de ITBMS)" not in filas[2][1]
    # el grupo va en negrita y sombreado
    celda_grupo = tabla_items(doc).rows[1].cells[1]
    assert celda_grupo.paragraphs[0].runs[0].bold is True
    t = textos(doc)
    assert "Administración (10%):" in t and "B/. 35.00" in t and "B/. 23.45" in t and "B/. 408.45" in t


def test_estado_oculto_por_defecto_y_visible_con_opcion(conn, tmp_path):
    word.exportar_cotizacion(1, str(tmp_path / "a.docx"))
    assert "Estado" not in textos(abrir(tmp_path / "a.docx"))
    db.guardar_config("mostrar_estado_pdf", "1", conn)
    word.exportar_cotizacion(1, str(tmp_path / "b.docx"))
    assert "Estado: borrador" in textos(abrir(tmp_path / "b.docx"))
    word.exportar_cotizacion(1, str(tmp_path / "c.docx"), mostrar_estado=False)
    assert "Estado" not in textos(abrir(tmp_path / "c.docx"))


def test_notas_con_vinetas_reales_y_secciones_en_negrita(obra, tmp_path):
    ruta = tmp_path / "c.docx"
    word.exportar_cotizacion(1, str(ruta))
    ps = abrir(ruta).paragraphs
    i = next(n for n, p in enumerate(ps) if p.text == "Alcance, exclusiones y condiciones")
    titulo = ps[i]
    assert titulo.runs[0].bold and titulo.paragraph_format.keep_with_next
    resto = ps[i + 1:]
    assert [(p.text, p.style.name) for p in resto] == [
        ("ALCANCE", "Normal"), ("Obra civil según planos.", "List Bullet"), ("Segunda viñeta", "List Bullet"),
        ("FORMA DE PAGO", "Normal"), ("50% al inicio.", "List Bullet"), ("Texto suelto", "Normal")]
    assert resto[0].runs[0].bold and resto[0].paragraph_format.keep_with_next


def test_sin_notas_no_hay_seccion(conn, tmp_path):
    word.exportar_cotizacion(1, str(tmp_path / "c.docx"))
    assert "Alcance, exclusiones" not in textos(abrir(tmp_path / "c.docx"))


def test_notas_largas_se_exportan_completas(conn, tmp_path):
    db.guardar_notas(1, "\n".join(f"• Condición número {i}" for i in range(200)), conn)
    word.exportar_cotizacion(1, str(tmp_path / "c.docx"))
    t = textos(abrir(tmp_path / "c.docx"))
    assert "Condición número 0" in t and "Condición número 199" in t


def test_factura_sin_tabla_de_items_con_totales_y_origen(obra, tmp_path):
    db.cambiar_estado_cotizacion(1, "enviada", obra)
    db.cambiar_estado_cotizacion(1, "aprobada", obra)
    db.generar_factura(1, obra)
    db.marcar_factura_pagada(1, "2026-02-01", obra)
    ruta = tmp_path / "f.docx"
    word.exportar_factura(1, str(ruta))
    doc = abrir(ruta)
    t = textos(doc)
    assert "FACTURA" in t and "FAC-2026-0001" in t and "Cotización de origen: COT-2026-0001" in t
    assert "Fecha de pago: 2026-02-01" in t and "B/. 408.45" in t and "Administración (10%):" in t
    assert not any(tb.rows[0].cells[0].text == "Ítem" for tb in doc.tables)
    assert "Términos de pago" in t and "Banco General" in t     # la factura nace con la plantilla de pago


def test_logo_de_la_empresa(conn, tmp_path, qapp):
    from PySide6.QtGui import QColor, QImage
    from app.ui_empresas import preparar_logo
    img = QImage(400, 200, QImage.Format_ARGB32)
    img.fill(QColor("#1D4ED8"))
    img.save(str(tmp_path / "logo.png"))
    eid = db.empresa_predeterminada_id(conn)
    db.guardar_empresa({"nombre": "Con Logo", "logo": preparar_logo(str(tmp_path / "logo.png"))}, conn, eid)
    word.exportar_cotizacion(1, str(tmp_path / "c.docx"))
    assert len(abrir(tmp_path / "c.docx").inline_shapes) == 1
    db.guardar_empresa({"nombre": "Logo roto", "logo": b"basura"}, conn, eid)
    word.exportar_cotizacion(1, str(tmp_path / "d.docx"))         # no debe fallar
    assert len(abrir(tmp_path / "d.docx").inline_shapes) == 0


def test_datos_de_la_empresa_en_el_encabezado(conn, tmp_path):
    db.guardar_empresa({"nombre": "Constructora Beta", "subtitulo": "OBRAS CIVILES", "ruc": "99-88",
                        "direccion": "Calle 50", "telefono": "222-3333", "correo": "a@b.com"},
                       conn, db.empresa_predeterminada_id(conn))
    word.exportar_cotizacion(1, str(tmp_path / "c.docx"))
    t = textos(abrir(tmp_path / "c.docx"))
    for esperado in ("Constructora Beta", "OBRAS CIVILES", "RUC: 99-88", "Calle 50", "Tel.: 222-3333", "a@b.com"):
        assert esperado in t, esperado
    assert "INTEGRO SA" not in t


def test_xml_respeta_el_orden_del_esquema(obra, tmp_path):
    """Word marca como dañado un .docx cuyos hijos de tcPr/tblPr estén desordenados."""
    word.exportar_cotizacion(1, str(tmp_path / "c.docx"))
    doc = abrir(tmp_path / "c.docx")
    nombre = lambda e: e.tag.split("}")[1]
    for tabla in doc.tables:
        indices = [word.ORDEN_TBL_PR.index(nombre(e)) for e in tabla._tbl.tblPr]
        assert indices == sorted(indices)
        for fila in tabla.rows:
            for celda in fila.cells:
                tc_pr = celda._tc.tcPr
                if tc_pr is not None:
                    indices = [word.ORDEN_TC_PR.index(nombre(e)) for e in tc_pr]
                    assert indices == sorted(indices)


def test_pagina_carta_con_margenes(conn, tmp_path):
    word.exportar_cotizacion(1, str(tmp_path / "c.docx"))
    s = abrir(tmp_path / "c.docx").sections[0]
    assert round(s.page_width.cm, 1) == 21.6 and round(s.page_height.cm, 1) == 27.9
    assert round(s.left_margin.cm, 1) == 2.0


def test_documentos_inexistentes(conn, tmp_path):
    with pytest.raises(ValueError):
        word.exportar_cotizacion(999, str(tmp_path / "x.docx"))
    with pytest.raises(ValueError):
        word.exportar_factura(999, str(tmp_path / "x.docx"))


def test_etiquetas_de_totales_compartidas(obra):
    t = db.totales_cotizacion(1, obra)
    etiquetas = documento.etiquetas_totales(t)
    assert [e[0] for e in etiquetas] == ["Subtotal:", "Administración (10%):", "ITBMS (7%):", "Total:"]
    assert etiquetas[-1] == ("Total:", t.total, True)


# ----------------------------------------------------------------- interfaz
def test_botones_exportar_word_en_cotizaciones_y_facturas(obra, qapp, monkeypatch, tmp_path):
    destino = tmp_path / "salida.docx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(destino), "")))
    infos = []
    monkeypatch.setattr(QMessageBox, "information", staticmethod(lambda *a: infos.append(a[2])))
    w = ui_main.MainWindow()
    tab = w.cotizaciones_tab
    tab.refrescar()
    tab.tabla.setCurrentCell(0, 0)
    assert tab.btn_word.isEnabled()
    tab.exportar_word()
    assert destino.exists() and "COT-2026-0001" in textos(abrir(destino)) and "salida.docx" in infos[-1]
    db.cambiar_estado_cotizacion(1, "enviada", obra)
    db.cambiar_estado_cotizacion(1, "aprobada", obra)
    db.generar_factura(1, obra)
    f = w.facturas_tab
    f.refrescar()
    f.tabla.setCurrentCell(0, 0)
    destino.unlink()
    f.exportar_word()
    assert "FAC-2026-0001" in textos(abrir(destino))


def test_exportar_word_cancelado_o_con_error(obra, qapp, monkeypatch, tmp_path):
    w = ui_main.MainWindow()
    tab = w.cotizaciones_tab
    tab.refrescar()
    tab.tabla.setCurrentCell(0, 0)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: ("", "")))
    tab.exportar_word()                                            # cancelar: no hace nada
    avisos = []
    monkeypatch.setattr(QMessageBox, "warning", staticmethod(lambda *a: avisos.append(a[2])))
    monkeypatch.setattr(QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: (str(tmp_path / "no_existe" / "x.docx"), "")))
    tab.exportar_word()                                            # carpeta inexistente: avisa
    assert avisos and "No se pudo exportar" in avisos[0]
