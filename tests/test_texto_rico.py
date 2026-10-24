import pytest
from docx import Document
from PySide6.QtGui import QTextListFormat

from app import db, pdf, texto_rico, ui_items, ui_main, word
from app.calculo import Linea
from app.ui_texto import EditorTextoDialog

RICO = "<p>Muro <b>reforzado</b> y <i>pintado</i></p><ul><li>uno</li><li>dos</li></ul><ol><li>a</li><li>b</li></ol>"


def test_ida_y_vuelta_y_texto_plano():
    assert texto_rico.serializar(texto_rico.parsear(RICO)) == RICO
    assert texto_rico.a_texto_plano(RICO) == "Muro reforzado y pintado\n• uno\n• dos\n1. a\n2. b"
    assert texto_rico.a_texto_plano("simple\nlinea") == "simple\nlinea"      # texto antiguo, sin cambios
    assert texto_rico.es_vacio("") and texto_rico.es_vacio("<p> </p>") and not texto_rico.es_vacio("x")


def test_editor_carga_y_guarda(qapp):
    d = EditorTextoDialog(RICO)
    assert d.texto() == RICO
    d = EditorTextoDialog("a\nb")
    c = d.editor.textCursor()
    c.select(c.SelectionType.Document)
    d.editor.setTextCursor(c)
    d._negrita()
    d._lista(QTextListFormat.ListDisc)
    assert d.texto() == "<ul><li><b>a</b></li><li><b>b</b></li></ul>"
    assert EditorTextoDialog("  ").texto() == ""


def test_texto_largo_no_se_pierde(qapp):
    largo = "<p>" + " ".join(f"palabra{i}" for i in range(3000)) + "</p>"
    assert texto_rico.a_texto_plano(EditorTextoDialog(largo).texto()).count("palabra") == 3000


@pytest.fixture
def cot(conn):
    db.guardar_cotizacion(1, [Linea("item", 0, RICO, 1, 100.0, True, "gl")], [], conn)
    return conn


def test_pdf_y_word_con_formato(cot, tmp_path, pdf_sin_comprimir):
    pdf.exportar_cotizacion(1, str(tmp_path / "c.pdf"))
    assert b"reforzado" in (tmp_path / "c.pdf").read_bytes()
    word.exportar_cotizacion(1, str(tmp_path / "c.docx"))
    doc = Document(str(tmp_path / "c.docx"))
    celda = next(t for t in doc.tables if t.rows[0].cells[0].text == "Ítem").rows[1].cells[1]
    assert "reforzado" in celda.text and "•\tuno" in celda.text and "1.\ta" in celda.text
    runs = [r for p in celda.paragraphs for r in p.runs]
    assert any(r.bold and r.text == "reforzado" for r in runs) and any(r.italic and r.text == "pintado" for r in runs)


def test_dialogo_items_edita_descripcion(cot, qapp, monkeypatch):
    d = ui_items.ItemsDialog(1, "COT-2026-0001", editable=True)
    monkeypatch.setattr(EditorTextoDialog, "exec", lambda self: self.editor.setHtml("<p>Nuevo <b>texto</b></p>") or 1)
    d.editar_descripcion(0)
    assert d.lineas[0].descripcion == "<p>Nuevo <b>texto</b></p>"
    assert d.tabla.item(0, ui_items.COL_DESC).text() == "Nuevo texto"


def test_titulo_notas_factura_y_cotizacion(cot, tmp_path, pdf_sin_comprimir):
    db.guardar_notas(1, "• Una condición", cot)
    assert db.obtener_titulo_notas("cotizaciones", 1, cot) == ""
    db.guardar_titulo_notas("cotizaciones", 1, "Condiciones de pago", cot)
    db.cambiar_estado_cotizacion(1, "enviada", cot)
    db.cambiar_estado_cotizacion(1, "aprobada", cot)
    num = db.generar_factura(1, cot)
    fid = cot.execute("SELECT id FROM facturas WHERE numero = ?", (num,)).fetchone()["id"]
    assert db.obtener_titulo_notas("facturas", fid, cot) == "Condiciones de pago"      # se copia a la factura
    db.guardar_titulo_notas("facturas", fid, "Términos de la factura", cot)
    pdf.exportar_factura(fid, str(tmp_path / "f.pdf"))
    assert b"T\xc3\xa9rminos de la factura" in (tmp_path / "f.pdf").read_bytes() or b"rminos de la factura" in (tmp_path / "f.pdf").read_bytes()
    word.exportar_factura(fid, str(tmp_path / "f.docx"))
    assert any(p.text == "Términos de la factura" for p in Document(str(tmp_path / "f.docx")).paragraphs)
    db.guardar_titulo_notas("facturas", fid, "", cot)
    word.exportar_factura(fid, str(tmp_path / "g.docx"))
    assert any(p.text == db.TITULO_NOTAS for p in Document(str(tmp_path / "g.docx")).paragraphs)


def test_dialogo_notas_con_titulo(qapp):
    d = ui_main.NotasDialog("x", "texto", None, "Mi título")
    assert d.titulo_bloque() == "Mi título"


def test_ocultar_texto_en_tabla(cot, qapp):
    d = ui_items.ItemsDialog(1, "COT-2026-0001", editable=True)
    celda = lambda: d.tabla.item(0, ui_items.COL_DESC)
    assert "uno" in celda().text()
    d.btn_ver_texto.setChecked(False)
    assert celda().text() == "Muro reforzado y pintado …" and "uno" in celda().toolTip()
    d.btn_ver_texto.setChecked(True)
    assert "uno" in celda().text()


def test_contraer_costos_y_notas(cot, qapp):
    d = ui_items.ItemsDialog(1, "COT-2026-0001", editable=True)
    d.show()
    assert d.tabla_cargos.isVisible() and d.cuerpo_notas.isVisible()
    d.btn_ver_cargos.setChecked(False)
    d.btn_ver_notas.setChecked(False)
    assert not d.tabla_cargos.isVisible() and not d.cuerpo_notas.isVisible()
    assert d.btn_ver_cargos.text() == "Expandir"
    d.btn_ver_notas.setChecked(True)
    assert d.cuerpo_notas.isVisible() and d.btn_plantilla.isVisible()


def test_ventana_de_cotizacion_se_puede_maximizar(cot, qapp):
    from PySide6.QtCore import Qt
    d = ui_items.ItemsDialog(1, "COT-2026-0001", editable=True)
    assert d.windowFlags() & Qt.WindowMaximizeButtonHint
    d.show()
    d.showMaximized()
    assert d.isMaximized() or d.windowState() & Qt.WindowMaximized
