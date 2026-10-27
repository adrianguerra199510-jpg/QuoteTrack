"""Exportación de cotizaciones y facturas a Word (.docx) con python-docx: mismo contenido que el PDF, editable."""
import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app import calculo, documento, texto_rico
from app.db import TITULO_NOTAS

FUENTE = "Arial"
GRIS = "D9D9D9"
GRIS_CLARO = "F2F2F2"
ANCHO_UTIL = 17.5  # cm (carta con márgenes de 2 cm)


def _dinero(valor: float) -> str:
    return f"B/. {valor:,.2f}"


# Orden de los hijos que exige el esquema de Word (un orden distinto hace que Word diga que el archivo está dañado)
ORDEN_TC_PR = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark"]
ORDEN_TBL_PR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
                "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
                "tblCaption", "tblDescription"]


def _insertar_ordenado(padre, hijo, orden) -> None:
    """Inserta `hijo` en `padre` respetando la secuencia `orden` del esquema OOXML."""
    nombre = lambda e: e.tag.split("}")[1]
    posicion = orden.index(nombre(hijo))
    for existente in list(padre):
        if nombre(existente) in orden and orden.index(nombre(existente)) > posicion:
            existente.addprevious(hijo)
            return
    padre.append(hijo)


def _sombrear(celda, color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    _insertar_ordenado(celda._tc.get_or_add_tcPr(), shd, ORDEN_TC_PR)


def _borde_superior(celda, grosor: int = 8) -> None:
    bordes = OxmlElement("w:tcBorders")
    b = OxmlElement("w:top")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(grosor))
    b.set(qn("w:color"), "000000")
    bordes.append(b)
    _insertar_ordenado(celda._tc.get_or_add_tcPr(), bordes, ORDEN_TC_PR)


def _sin_bordes(tabla) -> None:
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "nil")
        bordes.append(e)
    _insertar_ordenado(tabla._tbl.tblPr, bordes, ORDEN_TBL_PR)


def _anchos(tabla, anchos_cm) -> None:
    tabla.autofit = False
    for col, ancho in zip(tabla.columns, anchos_cm):
        col.width = Cm(ancho)             # cuadrícula de la tabla (la respetan Word y LibreOffice)
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos_cm):
            celda.width = Cm(ancho)


def _repetir_encabezado(fila) -> None:
    tr_pr = fila._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    tr_pr.append(e)


def _no_partir(fila) -> None:
    tr_pr = fila._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _texto(parrafo, texto, negrita=False, cursiva=False, tam=None, color=None):
    r = parrafo.add_run(texto)
    r.bold, r.italic = negrita, cursiva
    if tam:
        r.font.size = Pt(tam)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _parrafo(contenedor, texto="", alinear=None, negrita=False, cursiva=False, tam=None, antes=0, despues=2):
    p = contenedor.add_paragraph()
    if alinear is not None:
        p.alignment = alinear
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(despues)
    if texto:
        _texto(p, texto, negrita, cursiva, tam)
    return p


def _celda_texto(celda, texto, alinear=WD_ALIGN_PARAGRAPH.LEFT, negrita=False, sangria_cm=0.0):
    """Escribe en una celda (reutilizando su primer párrafo) un texto con formato: negrita, cursiva, subrayado y listas."""
    parrafos = texto_rico.parsear(texto) or [texto_rico.Parrafo("p", [])]
    for n, par in enumerate(parrafos):
        p = celda.paragraphs[0] if n == 0 else celda.add_paragraph()
        p.alignment = alinear
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(1)
        sangria = sangria_cm
        if par.tipo != "p":
            sangria += 0.5
            p.paragraph_format.first_line_indent = Cm(-0.4)
        if sangria:
            p.paragraph_format.left_indent = Cm(sangria)
        if par.tipo != "p":
            _texto(p, ("•" if par.tipo == "ul" else f"{par.numero}.") + "\t", negrita)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(sangria))
        for t in par.tramos:
            r = _texto(p, t.texto.replace("\n", "\v"), negrita or t.negrita, t.cursiva)
            r.underline = t.subrayado


def _encabezado_empresa(doc, empresa) -> None:
    if empresa is None:
        return
    t = doc.add_table(rows=1, cols=2)
    _sin_bordes(t)
    _anchos(t, [ANCHO_UTIL - 9, 9])
    izq, der = t.rows[0].cells
    if empresa["logo"]:
        try:
            izq.paragraphs[0].add_run().add_picture(io.BytesIO(empresa["logo"]), height=Cm(2.2))
        except Exception:  # logo corrupto: se exporta sin logo
            pass
    datos = [(empresa["nombre"], True, 16), (empresa["subtitulo"], False, 10),
             (f"RUC: {empresa['ruc']}" if empresa["ruc"] else "", False, 9), (empresa["direccion"], False, 9),
             (f"Tel.: {empresa['telefono']}" if empresa["telefono"] else "", False, 9), (empresa["correo"], False, 9)]
    primero = True
    for texto, negrita, tam in datos:
        if not (texto or "").strip():
            continue
        p = der.paragraphs[0] if primero else der.add_paragraph()
        primero = False
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(2)
        _texto(p, texto.strip(), negrita, tam=tam)
    _parrafo(doc, despues=4)


def _datos_cliente(doc, d: documento.Documento) -> None:
    _parrafo(doc, d.titulo, WD_ALIGN_PARAGRAPH.CENTER, negrita=True, tam=20, antes=6, despues=10)
    for etiqueta, valor in [("Número", d.numero), ("Fecha", d.fecha), *d.extra]:
        p = _parrafo(doc, despues=1)
        _texto(p, f"{etiqueta}: ", negrita=True)
        _texto(p, str(valor))
    _parrafo(doc, "Cliente", negrita=True, cursiva=True, tam=11, antes=12, despues=2).paragraph_format.keep_with_next = True
    cl = d.cliente
    _parrafo(doc, cl["nombre"], despues=1)
    if cl["identificacion"]:
        _parrafo(doc, f"RUC/Cédula: {cl['identificacion']}", despues=1)
    if cl["contacto"]:
        _parrafo(doc, f"Contacto: {cl['contacto']}", despues=1)
    _parrafo(doc, despues=4)


def _tabla_items(doc, lineas, tasa) -> None:
    cods, imps = calculo.codigos(lineas), calculo.importes(lineas)
    hay_exentos = tasa > 0 and any(not l.es_grupo and not l.aplica_itbms for l in lineas)
    anchos = [1.4, 6.2, 1.4, 2.1, 3.2, 3.2]
    t = doc.add_table(rows=1, cols=6)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    der, cen = WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER
    for celda, texto, al in zip(t.rows[0].cells, ["Ítem", "Descripción", "Ud.", "Cantidad", "Precio unitario", "Importe"],
                                [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, cen, der, der, der]):
        _celda_texto(celda, texto, al, negrita=True)
        _sombrear(celda, GRIS)
    _repetir_encabezado(t.rows[0])
    for i, l in enumerate(lineas):
        fila = t.add_row()
        _no_partir(fila)
        c = fila.cells
        if l.es_grupo:
            _celda_texto(c[0], cods[i], negrita=True)
            _celda_texto(c[1], texto_rico.con_titulo(l.titulo, l.descripcion), negrita=True)
            _celda_texto(c[5], _dinero(imps[i]), der, negrita=True)
            for celda in c:
                _sombrear(celda, GRIS_CLARO)
            continue
        _celda_texto(c[0], cods[i])
        _celda_texto(c[1], texto_rico.con_titulo(l.titulo, l.descripcion), sangria_cm=0.25 if l.nivel == 1 else 0.0)
        if hay_exentos and not l.aplica_itbms:
            _texto(c[1].paragraphs[-1], " (exento de ITBMS)", tam=7, color="666666")
        _celda_texto(c[2], l.unidad, cen)
        _celda_texto(c[3], f"{l.cantidad:g}", der)
        _celda_texto(c[4], _dinero(l.precio_mostrado), der)
        _celda_texto(c[5], _dinero(imps[i]), der)
    _anchos(t, anchos)
    _parrafo(doc, despues=4)


def _tabla_totales(doc, totales) -> None:
    filas = documento.etiquetas_totales(totales)
    t = doc.add_table(rows=len(filas), cols=2)
    _sin_bordes(t)
    t.alignment = WD_TABLE_ALIGNMENT.RIGHT
    for fila, (etiqueta, monto, es_total) in zip(t.rows, filas):
        _celda_texto(fila.cells[0], etiqueta, WD_ALIGN_PARAGRAPH.RIGHT, negrita=es_total)
        _celda_texto(fila.cells[1], _dinero(monto), WD_ALIGN_PARAGRAPH.RIGHT, negrita=es_total)
        if es_total:
            for celda in fila.cells:
                _borde_superior(celda)
    _anchos(t, [7.0, 3.8])


def _bloque_etapa(doc, d) -> None:
    """Facturas por etapas: concepto y resumen de lo facturado de la cotización."""
    if d.concepto:
        p = _parrafo(doc, "", antes=6, despues=6)
        _texto(p, "Concepto: ", negrita=True)
        _texto(p, d.concepto)
    if d.resumen:
        t = doc.add_table(rows=len(d.resumen), cols=2)
        _sin_bordes(t)
        t.alignment = WD_TABLE_ALIGNMENT.RIGHT
        for fila, (etiqueta, monto, ultimo) in zip(t.rows, d.resumen):
            _celda_texto(fila.cells[0], etiqueta, WD_ALIGN_PARAGRAPH.RIGHT, negrita=ultimo)
            _celda_texto(fila.cells[1], _dinero(monto), WD_ALIGN_PARAGRAPH.RIGHT, negrita=ultimo)
        _anchos(t, [7.0, 3.8])
        _parrafo(doc, "", despues=4)


def _notas(doc, notas, titulo="") -> None:
    """Título y notas con viñetas reales de Word (editables); las líneas en mayúsculas son encabezados de sección."""
    if not (notas or "").strip():
        return
    p = _parrafo(doc, (titulo or "").strip() or TITULO_NOTAS, negrita=True, cursiva=True, tam=11, antes=18, despues=4)
    p.paragraph_format.keep_with_next = True
    for linea in notas.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        texto = linea.strip()
        if not texto:
            continue
        if texto[0] in "•-" and len(texto) > 1 and texto[1] in " \t":
            sangrada = len(linea) - len(linea.lstrip()) >= 2
            p = doc.add_paragraph(texto[1:].strip(), style="List Bullet 2" if sangrada else "List Bullet")
            p.paragraph_format.space_after = Pt(1)
        elif texto.isupper():
            p = _parrafo(doc, texto, negrita=True, antes=8, despues=2)
            p.paragraph_format.keep_with_next = True
        else:
            _parrafo(doc, texto, despues=1)


def _construir(d: documento.Documento, ruta: str) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)        # carta
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FUENTE, Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    doc.styles["List Bullet"].font.name = doc.styles["List Bullet 2"].font.name = FUENTE
    doc.core_properties.title = f"{d.titulo.title()} {d.numero}"
    doc.core_properties.author = d.empresa["nombre"] if d.empresa else ""

    _encabezado_empresa(doc, d.empresa)
    _datos_cliente(doc, d)
    if d.lineas is not None:
        _tabla_items(doc, d.lineas, d.totales.tasa)
    _bloque_etapa(doc, d)
    _tabla_totales(doc, d.totales)
    _notas(doc, d.notas, d.titulo_notas)
    doc.save(ruta)


def exportar_cotizacion(cotizacion_id: int, ruta: str, mostrar_estado: bool | None = None) -> None:
    _construir(documento.de_cotizacion(cotizacion_id, mostrar_estado), ruta)


def exportar_factura(factura_id: int, ruta: str, mostrar_estado: bool | None = None) -> None:
    _construir(documento.de_factura(factura_id, mostrar_estado), ruta)
